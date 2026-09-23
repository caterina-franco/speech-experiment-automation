import pandas as pd
import os, random
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# --- 1. CONFIGURAZIONE DATI ---
LISTA_NOMI = ["Speaker_A", "Speaker_B", "Speaker_C", "Speaker_D"]

DB_VIRTUALI = {
    "Speaker_A":  ["Food", "Language", "Animal", "Hobby"],
    "Speaker_C": ["Place", "Entertainment", "Animal", "Food"],
    "Speaker_B": ["Hobby", "Place", "Subject", "Language"],
    "Speaker_D":   ["Food", "Hobby", "Language", "Superpower"]
}

DB_REALI = {
    "Speaker_A":   ["Music", "Travel", "Job", "Subject", "Entertainment", "Superpower"],
    "Speaker_C":  ["Music", "Hobby", "Job", "Subject", "Language", "Superpower"],
    "Speaker_B":  ["Music", "Food", "Job", "Entertainment", "Animals", "Superpower"],
    "Speaker_D":    ["Music", "Travel", "Job", "Subject", "Entertainment", "Animals"]
}

def crea_esperimento():
    lista_finale = []
    # A. Creazione base dei 32 trial
    for nome in LISTA_NOMI:
        for arg in DB_VIRTUALI[nome]:
            lista_finale.append({'Persona': nome, 'Argomento': arg, 'Modalità': 'VIRTUAL'})
        argomenti_reali = random.sample(DB_REALI[nome], 4)
        for arg in argomenti_reali:
            lista_finale.append({'Persona': nome, 'Argomento': arg, 'Modalità': 'REALE'})

    # B. Bilanciamento Altoparlanti 
    posizioni_R = [1, 2, 3, 4] * 4
    posizioni_V = [1, 2, 3, 4] * 4
    random.shuffle(posizioni_R)
    random.shuffle(posizioni_V)

    # C. Assegnazione posizioni
    random.shuffle(lista_finale)
    for trial in lista_finale:
        n_altoparlante = posizioni_R.pop() if trial['Modalità'] == 'REALE' else posizioni_V.pop()
        trial['Altoparlante'] = n_altoparlante
        
        altri_nomi = [n for n in LISTA_NOMI if n != trial['Persona']]
        random.shuffle(altri_nomi)
        ordine_visivo = [None] * 4
        ordine_visivo[n_altoparlante - 1] = trial['Persona']
        
        it = iter(altri_nomi)
        testo_posizioni = []
        for i in range(4):
            nome_da_inserire = ordine_visivo[i] if ordine_visivo[i] else next(it)
            testo_posizioni.append(f"Pos {i+1}: {nome_da_inserire}")
        trial['Posizioni'] = "\n".join(testo_posizioni)

    random.shuffle(lista_finale)
    for i, trial in enumerate(lista_finale, 1):
        trial['Trial'] = i
    return lista_finale

# --- 2. GENERAZIONE EXCEL ---
NUM_PARTECIPANTI = 6
NOME_FILE = "NOME.xlsx"
COLORE_VIOLA = "7644A6"
STILE_VIOLA = PatternFill("solid", start_color=COLORE_VIOLA)
STILE_VIRTUAL = PatternFill("solid", start_color="F2EBF5")
BORDI = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
CENTRATO = Alignment(horizontal='center', vertical='center', wrap_text=True)

with pd.ExcelWriter(NOME_FILE, engine='openpyxl') as writer:
    for n in range(16, 16 + NUM_PARTECIPANTI):
        id_str = f"{n:02d}"
        nome_foglio = f"Partecipante {id_str}"
        dati_trial = crea_esperimento()
        
        df_dati = pd.DataFrame(dati_trial)[['Trial', 'Posizioni', 'Persona', 'Modalità', 'Altoparlante', 'Argomento']]
        
        # Aggiunta la colonna 'Valutazione sicurezza' nel dizionario
        df_risposte = pd.DataFrame({
            'SÌ': ['']*32, 
            'NO': ['']*32, 
            'Valutazione sicurezza': ['']*32, 
            'Note': ['']*32
        })

        df_dati.to_excel(writer, sheet_name=nome_foglio, index=False, startrow=5)
        df_risposte.to_excel(writer, sheet_name=nome_foglio, index=False, startrow=5, startcol=7)

        ws = writer.sheets[nome_foglio]
        
        # --- INTESTAZIONI ---
        ws['A1'] = f"ESPERIMENTO PARTECIPANTE: {id_str}"
        ws['A1'].font = Font(bold=True, size=18, color=COLORE_VIOLA)
        
        ws['A2'] = "PARTECIPANTE: ________________________________"
        ws['A2'].font = Font(bold=True)
        
        ws['H2'] = "La voce che hai appena sentito proviene da un altoparlante?"
        ws['H2'].font = Font(bold=True, size=14)
        ws['H3'] = "Quanto sei sicuro/a da 1 a 5?"
        ws['H3'].font = Font(bold=True, size=14)
        
        ws['H5'] = "REGISTRO RISPOSTE"
        ws['H5'].font = Font(bold=True, color=COLORE_VIOLA)

        # Formattazione: aggiunto l'indice 11 per coprire la nuova colonna (K)
        for i in range(33): # Intestazione tabella + 32 righe
            riga_excel = i + 6
            ws.row_dimensions[riga_excel].height = 65 if i > 0 else 30
            
            # Formattazione colonne: aggiunta la colonna 11
            for col_excel in [1, 2, 3, 4, 5, 6, 8, 9, 10, 11]:
                cella = ws.cell(row=riga_excel, column=col_excel)
                cella.border = BORDI
                cella.alignment = CENTRATO
                
                if i == 16:
                    # Applichiamo un bordo inferiore più spesso o doppio per marcare la fine della prima metà
                    cella.border = Border(
                        left=Side(style='thin'), 
                        right=Side(style='thin'), 
                        top=Side(style='thin'), 
                        bottom=Side(style='medium') # 'medium' o 'double' crea un distacco visivo
                    )
                if i == 0: # Header tabella
                    cella.font = Font(bold=True, color="FFFFFF")
                    cella.fill = STILE_VIOLA
                else: # Righe dati
                    if dati_trial[i-1]['Modalità'] == 'VIRTUAL':
                        cella.fill = STILE_VIRTUAL

        # Aggiornate le larghezze: aggiunta una misura per la colonna K (11-esima)
        larghezze = [8, 25, 15, 12, 12, 20, 5, 10, 10, 15, 25]
        for i, largh in enumerate(larghezze):
            ws.column_dimensions[chr(65 + i)].width = largh

print(f"Fatto! Creato file '{NOME_FILE}'.")
os.startfile(NOME_FILE) if os.name == 'nt' else os.system(f"open '{NOME_FILE}'")